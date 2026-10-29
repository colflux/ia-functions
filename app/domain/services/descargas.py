"""Excel con los datos que la persona consultó en el chat.

Cuando una herramienta de datos se llama con exportar=true, devuelve todas las
filas (no solo la muestra que ve el modelo). El orquestador las junta: cada
consulta es una hoja del mismo libro. El archivo se guarda en el bucket bajo
descargas/ y el chat muestra un botón que pide un enlace temporal. Los datos ya
son públicos en el geoportal, así que cualquiera puede descargarlo.

Si todas las hojas traen el mismo "nombre" (por ejemplo el sitio del grupo de
observación de donde son las fotos), el archivo se llama con él:
colflux-<nombre>-<fecha>.xlsx.

Una consulta puede traer hojas de apoyo ("hojas_extra", por ejemplo «Por
persona y foto» o «Preguntas» en los grupos de observación): van justo después
de su hoja, en el mismo libro. Las filas que se informan son las de las hojas
principales. Una hoja con "ajustar_texto" (textos largos, como la descripción
detallada de una foto) lleva columnas más anchas y el texto ajustado.
"""
import io
import logging
import re
import uuid
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.domain.ports.file_storage import FileStorage

logger = logging.getLogger("uvicorn.error")

PREFIJO = "descargas/"
SEGUNDOS_ENLACE = 3600
MAX_FILAS_HOJA = 100_000
TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _nombre_hoja(titulo: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", titulo or "Datos").strip()[:28] or "Datos"
    nombre, n = base, 2
    while nombre.lower() in usados:
        nombre, n = f"{base[:26]} {n}", n + 1
    usados.add(nombre.lower())
    return nombre


def _limpio(nombre: Any) -> str:
    """Solo letras minúsculas, números y guiones: va en la clave del bucket y en el nombre del archivo."""
    return re.sub(r"[^a-z0-9]+", "-", str(nombre or "").lower()).strip("-")[:40]


class Descargas:
    def __init__(self, almacen: FileStorage | None) -> None:
        self._almacen = almacen

    @property
    def disponible(self) -> bool:
        return self._almacen is not None

    def guardar(self, hojas: list[dict[str, Any]]) -> dict[str, Any] | None:
        """hojas: [{"titulo": str, "filas": [dict, ...], "hojas_extra": [...]}]. Devuelve {archivo, nombre, filas, hojas}."""
        if self._almacen is None or not hojas:
            return None
        libro = Workbook()
        libro.remove(libro.active)
        usados: set[str] = set()
        total = 0
        principales = {id(h) for h in hojas}
        todas = [x for h in hojas for x in [h, *(h.get("hojas_extra") or [])]]
        for hoja in todas:
            filas = (hoja.get("filas") or [])[:MAX_FILAS_HOJA]
            ws = libro.create_sheet(_nombre_hoja(hoja.get("titulo", "Datos"), usados))
            columnas: list[str] = []
            for f in filas:
                columnas.extend(k for k in f if k not in columnas)
            ws.append(columnas or ["sin datos"])
            for celda in ws[1]:
                celda.font = Font(bold=True)
            for f in filas:
                ws.append([f.get(c) if not isinstance(f.get(c), (dict, list)) else str(f.get(c)) for c in columnas])
            ws.freeze_panes = "A2"
            ajustar = bool(hoja.get("ajustar_texto"))
            for i, c in enumerate(columnas, 1):
                ancho = max([len(str(c))] + [len(str(f.get(c) or "")) for f in filas[:200]])
                ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 60 if ajustar else 45)
            if ajustar:
                for fila in ws.iter_rows(min_row=2):
                    for celda in fila:
                        celda.alignment = Alignment(wrap_text=True, vertical="top")
            if id(hoja) in principales:
                total += len(filas)
        contenido = io.BytesIO()
        libro.save(contenido)
        nombres = {_limpio(h.get("nombre")) for h in hojas}
        nombre = nombres.pop() if len(nombres) == 1 else ""
        sufijo = f"-{nombre}" if nombre else ""
        clave = f"{PREFIJO}{date.today():%Y-%m-%d}-{uuid.uuid4().hex[:8]}{sufijo}.xlsx"
        self._almacen.subir(clave, contenido.getvalue(), TIPO_XLSX, {"filas": str(total)})
        logger.info("EXCEL %s con %d hojas y %d filas", clave, len(todas), total)
        return {"archivo": clave, "nombre": f"colflux-{nombre or 'datos'}-{date.today():%Y-%m-%d}.xlsx",
                "filas": total, "hojas": [h.get("titulo") for h in todas]}

    def enlace(self, clave: str) -> str | None:
        if self._almacen is None or not clave.startswith(PREFIJO) or ".." in clave or not clave.endswith(".xlsx"):
            return None
        con_nombre = re.fullmatch(r"(\d{4}-\d{2}-\d{2})-[0-9a-f]{8}-([a-z0-9-]+)\.xlsx", clave[len(PREFIJO):])
        archivo = (f"colflux-{con_nombre.group(2)}-{con_nombre.group(1)}.xlsx" if con_nombre
                   else "colflux-" + clave[len(PREFIJO):])
        return self._almacen.enlace_descarga(clave, archivo, SEGUNDOS_ENLACE)
